# data_handler.py
import os
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from collector import get_data_file_path, HardwareCollector


class DataHandler:
    def __init__(self):

        self.data_dir = os.path.join("数据存储")
        self.json_path = get_data_file_path("computer_assets.json")

        if not os.path.exists(self.data_dir):
            os.makedirs(self.data_dir, exist_ok=True)

        # 这里的名称建议与 get_hardware_info 返回的 key 顺序尽量保持一致
        self.column_order = [
            "计算机名称",
            "ip地址",
            "MAC地址",
            "当前用户",
            "品牌",
            "型号",
            "设备类型",
            "序列号",
            "系统UUID",
            "主板序列号",
            "BIOS版本",
            "CPU",
            "内存",
            "内存详情",
            "硬盘",
            "硬盘详情",
            "硬盘健康",
            "显卡",
            "主板信息",
            "操作系统",
            "部门",
            "现使用人",
            "收集时间",
        ]

        # 表头样式（蓝底白字加粗）
        self.header_fill = PatternFill("solid", fgColor="005FB8")
        self.header_font = Font(color="FFFFFF", bold=True, size=11)
        self.header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        # 健康异常行高亮（浅红）
        self.warn_fill = PatternFill("solid", fgColor="FFC7CE")
        self.warn_font = Font(color="9C0006")
        # 单元格细边框
        self.thin = Side(style="thin", color="D9D9D9")
        self.border = Border(left=self.thin, right=self.thin, top=self.thin, bottom=self.thin)

    def load_data(self):
        if os.path.exists(self.json_path):
            try:
                with open(self.json_path, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except:
                return []
        return []

    def save_data(self, data_list):
        try:
            with open(self.json_path, 'w', encoding='utf-8') as f:
                json.dump(data_list, f, ensure_ascii=False, indent=4)
            return True
        except:
            return False

    # ------------------------------------------------------------------
    # 通用：把数据写成带样式的 Excel 工作表
    # ------------------------------------------------------------------
    def _write_sheet(self, ws, headers, rows):
        """写入表头 + 数据，并统一套用样式、冻结、列宽、筛选、健康高亮"""
        # 1. 表头
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.fill = self.header_fill
            c.font = self.header_font
            c.alignment = self.header_align
            c.border = self.border

        # 2. 数据行
        health_col = headers.index("硬盘健康") + 1 if "硬盘健康" in headers else None
        for ri, row in enumerate(rows, 2):
            for ci, val in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.alignment = Alignment(vertical="center", wrap_text=False)
                c.border = self.border
            # 健康异常高亮（异常/警告/失败/损坏）
            if health_col is not None:
                hv = str(row[health_col - 1])
                if any(k in hv for k in ("异常", "警告", "失败", "损坏", "坏")):
                    for ci in range(1, len(headers) + 1):
                        cell = ws.cell(row=ri, column=ci)
                        cell.fill = self.warn_fill
                        cell.font = self.warn_font

        # 3. 冻结首行 + 自动筛选
        ws.freeze_panes = "A2"
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{max(1, len(rows) + 1)}"

        # 4. 自动列宽（按内容，限 8~46）
        for ci, h in enumerate(headers, 1):
            max_len = len(str(h))
            for ri in range(2, len(rows) + 2):
                v = ws.cell(row=ri, column=ci).value
                if v is not None:
                    # 中文按 2 个字符宽度计，英文 1
                    w = sum(2 if ord(ch) > 127 else 1 for ch in str(v))
                    max_len = max(max_len, w)
            ws.column_dimensions[get_column_letter(ci)].width = max(8, min(46, max_len + 2))

    # ------------------------------------------------------------------
    # 统计工作表：按 品牌 / 设备类型 / 部门 计数
    # ------------------------------------------------------------------
    def _write_stats(self, ws, data):
        def count_by(key):
            d = {}
            for item in data:
                v = str(item.get(key, "未填写") or "未填写").strip() or "未填写"
                d[v] = d.get(v, 0) + 1
            return sorted(d.items(), key=lambda kv: (-kv[1], kv[0]))

        blocks = [("按品牌统计", "品牌"), ("按设备类型统计", "设备类型"), ("按部门统计", "部门")]
        row = 1
        for title, key in blocks:
            ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=12, color="005FB8")
            row += 1
            h1 = ws.cell(row=row, column=1, value=key)
            h2 = ws.cell(row=row, column=2, value="数量")
            for h in (h1, h2):
                h.fill = self.header_fill
                h.font = self.header_font
                h.alignment = self.header_align
                h.border = self.border
            row += 1
            for name, cnt in count_by(key):
                a = ws.cell(row=row, column=1, value=name)
                b = ws.cell(row=row, column=2, value=cnt)
                a.border = self.border
                b.border = self.border
                b.alignment = Alignment(horizontal="center")
                row += 1
            row += 1  # 块之间空一行
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 10

    # ------------------------------------------------------------------
    # 导出 Excel：资产清单（带样式）+ 统计 sheet
    # ------------------------------------------------------------------
    def export_to_excel(self, data):
        try:
            if not data:
                return False, "数据为空"

            export_dir = os.path.join(self.data_dir, "导出")
            os.makedirs(export_dir, exist_ok=True)

            wb = Workbook()
            ws = wb.active
            ws.title = "硬件资产清单"

            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

            # 表头：优先预定义列顺序，再补数据里出现但未定义的字段
            headers = list(self.column_order)
            extra_keys = set()
            for item in data:
                extra_keys.update(item.keys())
            for k in sorted(extra_keys):
                if k not in headers:
                    headers.append(k)

            rows = [[str(item.get(h, "")) for h in headers] for item in data]

            self._write_sheet(ws, headers, rows)

            # 第二个工作表：统计
            stats_ws = wb.create_sheet("统计")
            self._write_stats(stats_ws, data)

            export_path = os.path.join(export_dir, f"资产清单导出{timestamp}.xlsx")
            wb.save(export_path)
            return True, export_path
        except Exception as e:
            return False, str(e)
