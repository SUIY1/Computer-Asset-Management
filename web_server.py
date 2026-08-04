# web_server.py
# 网站版「中心采集服务器」（Flask）
# ------------------------------------------------------------
# 功能：
#   1) 接收各机器的 agent 上报（POST /api/report），落盘为 JSON
#   2) 提供网页仪表盘（GET /）：资产列表、搜索、健康着色、导出
#   3) 提供同步接口（GET /api/machines），与客户端「从服务器同步」互通
#   4) 导出 Excel / CSV（GET /export）
# 部署：gunicorn web_server:app -b 0.0.0.0:8000  （前面用 Nginx 反代即可）
# 注意：本文件刻意【不依赖任何 Windows 专用模块】，可直接跑在 Linux 云服务器上。
import os
import io
import csv
import json
from datetime import datetime

from flask import (
    Flask, request, jsonify, render_template_string, send_file, abort,
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REPORT_DIR = os.path.join(BASE_DIR, "数据存储", "批量上报")
os.makedirs(REPORT_DIR, exist_ok=True)

app = Flask(__name__)


# ------------------------- 存储（与桌面端 server.py 同目录结构，便于互通） -------------------------
def _safe_name(name):
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in str(name or "unknown"))


def list_reports():
    out = []
    try:
        names = sorted(os.listdir(REPORT_DIR))
    except Exception:
        return out
    for fn in names:
        if fn.lower().endswith(".json"):
            try:
                with open(os.path.join(REPORT_DIR, fn), "r", encoding="utf-8") as f:
                    out.append(json.load(f))
            except Exception:
                pass
    return out


def save_report(data):
    name = data.get("计算机名称") or data.get("hostname") or "unknown"
    path = os.path.join(REPORT_DIR, f"{_safe_name(name)}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return path


def apply_query(machines, q):
    """按关键词过滤（与仪表盘搜索逻辑一致）：任意字段包含关键词即命中"""
    if not q:
        return machines
    q = q.strip().lower()
    return [m for m in machines if any(q in str(v).lower() for v in m.values())]


# 仪表盘展示的列（key 对应采集字段）
COLUMNS = [
    ("计算机名称", "计算机名称"),
    ("品牌", "品牌"),
    ("型号", "型号"),
    ("设备类型", "设备类型"),
    ("序列号", "序列号"),
    ("系统UUID", "系统UUID"),
    ("CPU", "CPU"),
    ("内存", "内存"),
    ("内存详情", "内存详情"),
    ("硬盘", "硬盘"),
    ("硬盘健康", "硬盘健康"),
    ("显卡", "显卡"),
    ("主板信息", "主板信息"),
    ("ip地址", "ip地址"),
    ("MAC地址", "MAC地址"),
    ("操作系统", "操作系统"),
    ("部门", "部门"),
    ("现使用人", "现使用人"),
    ("上报时间", "上报时间"),
    ("来源", "来源"),
]


# ------------------------- 接口 -------------------------
@app.route("/api/report", methods=["POST", "OPTIONS"])
def api_report():
    if request.method == "OPTIONS":
        return ("", 204)
    try:
        data = request.get_json(force=True, silent=True)
    except Exception:
        return jsonify({"ok": False, "error": "invalid json"}), 400
    if not isinstance(data, dict):
        return jsonify({"ok": False, "error": "body must be a json object"}), 400
    path = save_report(data)
    return jsonify({"ok": True, "saved": os.path.basename(path)})


@app.route("/api/machines")
def api_machines():
    return jsonify(list_reports())


@app.route("/")
def dashboard():
    q = (request.args.get("q") or "").strip()
    machines = apply_query(list_reports(), q)
    return render_template_string(
        DASHBOARD_TPL,
        machines=machines,
        q=q,
        count=len(machines),
        now=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )


@app.route("/export")
def export():
    fmt = (request.args.get("fmt") or "xlsx").lower()
    q = (request.args.get("q") or "").strip()
    machines = apply_query(list_reports(), q)  # 与仪表盘搜索一致：导出当前筛选结果
    keys = [k for _, k in COLUMNS]

    if fmt == "csv":
        return export_csv_fallback(machines, keys)

    # 默认 Excel（带样式 + 统计页）
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception:
        return export_csv_fallback(machines, keys)

    header_fill = PatternFill("solid", fgColor="005FB8")
    header_font = Font(color="FFFFFF", bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    warn_fill = PatternFill("solid", fgColor="FFC7CE")
    warn_font = Font(color="9C0006")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "资产列表"

    # 表头
    for ci, k in enumerate(keys, 1):
        c = ws.cell(row=1, column=ci, value=k)
        c.fill = header_fill
        c.font = header_font
        c.alignment = header_align
        c.border = border

    # 数据行
    health_idx = keys.index("硬盘健康") + 1 if "硬盘健康" in keys else None
    for ri, m in enumerate(machines, 2):
        for ci, k in enumerate(keys, 1):
            c = ws.cell(row=ri, column=ci, value=str(m.get(k, "")))
            c.alignment = Alignment(vertical="center")
            c.border = border
        if health_idx is not None:
            hv = str(m.get("硬盘健康", ""))
            if any(x in hv for x in ("异常", "警告", "失败", "损坏", "坏")):
                for ci in range(1, len(keys) + 1):
                    cc = ws.cell(row=ri, column=ci)
                    cc.fill = warn_fill
                    cc.font = warn_font

    # 冻结首行 + 自动筛选
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(keys))
    ws.auto_filter.ref = f"A1:{last_col}{max(1, len(machines) + 1)}"

    # 自动列宽（中文按 2 计，限 8~46）
    for ci, k in enumerate(keys, 1):
        ml = len(str(k))
        for ri in range(2, len(machines) + 2):
            v = ws.cell(row=ri, column=ci).value
            if v is not None:
                w = sum(2 if ord(ch) > 127 else 1 for ch in str(v))
                ml = max(ml, w)
        ws.column_dimensions[get_column_letter(ci)].width = max(8, min(46, ml + 2))

    # 统计页：按 品牌 / 设备类型 / 部门
    stats = wb.create_sheet("统计")

    def count_by(key):
        d = {}
        for m in machines:
            v = str(m.get(key, "未填写") or "未填写").strip() or "未填写"
            d[v] = d.get(v, 0) + 1
        return sorted(d.items(), key=lambda kv: (-kv[1], kv[0]))

    row = 1
    for title, key in (("按品牌统计", "品牌"),
                       ("按设备类型统计", "设备类型"),
                       ("按部门统计", "部门")):
        stats.cell(row=row, column=1, value=title).font = Font(bold=True, size=12, color="005FB8")
        row += 1
        for j, h in enumerate(("类别", "数量"), 1):
            c = stats.cell(row=row, column=j, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = header_align
            c.border = border
        row += 1
        for name, cnt in count_by(key):
            a = stats.cell(row=row, column=1, value=name)
            b = stats.cell(row=row, column=2, value=cnt)
            a.border = border
            b.border = border
            b.alignment = Alignment(horizontal="center")
            row += 1
        row += 1
    stats.column_dimensions["A"].width = 22
    stats.column_dimensions["B"].width = 10

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(
        out,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"资产导出_{datetime.now():%Y%m%d}.xlsx",
    )


def export_csv_fallback(machines, keys):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(keys)
    for m in machines:
        w.writerow([str(m.get(k, "")) for k in keys])
    data = buf.getvalue().encode("utf-8-sig")
    return send_file(
        io.BytesIO(data),
        mimetype="text/csv",
        as_attachment=True,
        download_name=f"资产导出_{datetime.now():%Y%m%d}.csv",
    )


@app.after_request
def _cors(resp):
    resp.headers["Access-Control-Allow-Origin"] = "*"
    resp.headers["Access-Control-Allow-Headers"] = "Content-Type"
    resp.headers["Access-Control-Allow-Methods"] = "GET,POST,OPTIONS"
    return resp


# ------------------------- 仪表盘模板（内联 CSS，离线可用） -------------------------
DASHBOARD_TPL = """<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>计算机资产采集 · 仪表盘</title>
<style>
 body{margin:0;font-family:-apple-system,'Microsoft YaHei',Segoe UI,sans-serif;background:#f5f6f8;color:#222}
 header{background:#005fb8;color:#fff;padding:16px 24px;display:flex;align-items:center;justify-content:space-between;flex-wrap:wrap;gap:8px}
 header h1{font-size:18px;margin:0}
 .wrap{max-width:1280px;margin:0 auto;padding:18px 20px}
 .bar{display:flex;align-items:center;gap:10px;flex-wrap:wrap;margin-bottom:14px}
 .bar input{padding:8px 10px;border:1px solid #cfd6dd;border-radius:6px;min-width:260px;font-size:14px}
 .bar button,.bar a.btn{padding:8px 14px;border:none;border-radius:6px;background:#005fb8;color:#fff;cursor:pointer;font-size:14px;text-decoration:none}
 .bar a.ghost{background:#eef3f8;color:#005fb8;border:1px solid #cfe0f0}
 .count{color:#666;font-size:14px}
 .hint{background:#fff8e1;border:1px solid #ffe082;color:#7a5b00;padding:10px 14px;border-radius:8px;font-size:13px;margin-bottom:14px;line-height:1.6}
 table{border-collapse:collapse;width:100%;background:#fff;box-shadow:0 1px 3px rgba(0,0,0,.08);border-radius:8px;overflow:hidden}
 th,td{border-bottom:1px solid #eef1f4;padding:9px 10px;font-size:13px;text-align:left;max-width:260px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
 th{background:#eef3f8;position:sticky;top:0;font-weight:bold}
 tr:nth-child(even){background:#fafbfc}
 .ok{color:#2e7d32;font-weight:bold}
 .warn{color:#e65100;font-weight:bold}
 .bad{color:#c62828;font-weight:bold}
 .empty{padding:40px;text-align:center;color:#999}
 code{background:#eef3f8;padding:1px 6px;border-radius:4px;font-size:12px}
</style></head>
<body>
<header><h1>🖥️ 计算机资产采集仪表盘</h1>
 <div class="count">共 <b>{{ count }}</b> 台 · 刷新于 {{ now }}</div>
</header>
<div class="wrap">
 <div class="hint">
  📥 <b>如何让机器上报到这里？</b> 在桌面端「批量采集中心」设置本服务器地址（如 <code>http://你的服务器IP:端口</code>），
  生成部署包 → 目标机双击 <code>agent.exe</code> 即自动上报。也可命令行：<code>agent.exe --server http://地址:端口</code>。
  接收接口：<code>POST /api/report</code>　同步接口：<code>GET /api/machines</code>。
 </div>
 <div class="bar">
   <form method="get" style="display:flex;gap:10px;flex-wrap:wrap">
     <input name="q" value="{{ q }}" placeholder="搜索 计算机名/品牌/型号/IP/使用人…">
     <button type="submit">🔍 搜索</button>
   </form>
   <a class="btn" href="/export?fmt=xlsx{% if q %}&amp;q={{ q }}{% endif %}">⬇ 导出 Excel</a>
   <a class="btn ghost" href="/export?fmt=csv{% if q %}&amp;q={{ q }}{% endif %}">⬇ 导出 CSV</a>
 </div>
 {% if q %}
 <div class="count" style="margin:-4px 0 10px">当前已按关键词「<b>{{ q }}</b>」筛选，<b>导出将只含筛选结果</b>（共 {{ count }} 台）。清空搜索后导出全部。</div>
 {% endif %}
 {% if machines %}
 <div style="overflow:auto;max-height:72vh">
 <table>
  <thead><tr>{% for label,_ in columns %}<th>{{ label }}</th>{% endfor %}</tr></thead>
  <tbody>
  {% for m in machines %}
   <tr>{% for _,k in columns %}
     <td>
     {% if k=='硬盘健康' %}
        {% set h=m.get(k,'') %}
        {% if '不良' in h or '坏' in h %}<span class="bad">{{ h }}</span>
        {% elif '警告' in h or '需关注' in h %}<span class="warn">{{ h }}</span>
        {% elif h and h!='未知' %}<span class="ok">{{ h }}</span>
        {% else %}{{ h }}{% endif %}
     {% else %}{{ m.get(k,'') }}{% endif %}
     </td>
   {% endfor %}</tr>
  {% endfor %}
  </tbody>
 </table>
 </div>
 {% else %}
 <div class="empty">暂无上报数据。请先让目标机运行 agent 上报到本服务器。</div>
 {% endif %}
</div>
</body></html>
"""

# 把 COLUMNS 传给模板（render_template_string 默认不暴露模块变量，这里手动注入）
@app.context_processor
def inject_columns():
    return {"columns": COLUMNS}


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "8000"))
    print(f"资产采集网站版已启动：http://0.0.0.0:{port}  （Ctrl+C 停止）")
    app.run(host="0.0.0.0", port=port, debug=False)
