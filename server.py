# server.py
# 中心采集服务器：接收各机器上报的硬件信息(JSON)，并提供网页仪表盘查看。
# 仅使用 Python 标准库（http.server / ThreadingHTTPServer），无需额外依赖。
import json
import os
import threading
import webbrowser
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs, quote

from collector import get_data_file_path

# 上报数据存放目录：数据存储/批量上报
REPORT_DIR = get_data_file_path("批量上报")


def ensure_report_dir():
    os.makedirs(REPORT_DIR, exist_ok=True)


def _safe_name(name):
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in str(name or "unknown"))


def list_reports():
    """读取所有已上报的机器信息"""
    ensure_report_dir()
    out = []
    for fn in sorted(os.listdir(REPORT_DIR)):
        if fn.lower().endswith(".json"):
            try:
                with open(os.path.join(REPORT_DIR, fn), "r", encoding="utf-8") as f:
                    out.append(json.load(f))
            except Exception:
                pass
    return out


def save_report(data):
    """保存一份上报数据为 <计算机名称>.json（重复上报会覆盖更新）"""
    ensure_report_dir()
    name = data.get("计算机名称") or data.get("hostname") or "unknown"
    path = os.path.join(REPORT_DIR, f"{_safe_name(name)}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return path


def apply_query(machines, q):
    """按关键词过滤（与网页版一致）：任意字段包含关键词即命中"""
    if not q:
        return machines
    q = q.strip().lower()
    return [m for m in machines if any(q in str(v).lower() for v in m.values())]


# 导出用的列（与网页版 web_server / 桌面端导出保持一致）
EXPORT_COLS = [
    ("计算机名称", "计算机名称"),
    ("品牌", "品牌"),
    ("型号", "型号"),
    ("设备类型", "设备类型"),
    ("序列号", "序列号"),
    ("系统UUID", "系统UUID"),
    ("CPU", "CPU"),
    ("内存", "内存"),
    ("内存详情", "内存详情"),
    ("硬盘", "硬盘"),
    ("硬盘健康", "硬盘健康"),
    ("显卡", "显卡"),
    ("主板信息", "主板信息"),
    ("ip地址", "ip地址"),
    ("MAC地址", "MAC地址"),
    ("操作系统", "操作系统"),
    ("部门", "部门"),
    ("现使用人", "现使用人"),
    ("来源", "来源"),
    ("上报时间", "上报时间"),
]


def build_export(machines, fmt):
    """返回 (content_type, bytes, filename)。fmt=xlsx 用 openpyxl 带样式+统计页；否则 CSV。"""
    import io
    import csv
    from datetime import datetime

    keys = [k for _, k in EXPORT_COLS]
    date = datetime.now().strftime("%Y%m%d")

    def _csv_bytes():
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(keys)
        for m in machines:
            w.writerow([str(m.get(k, "")) for k in keys])
        return ("\ufeff" + buf.getvalue()).encode("utf-8")

    if fmt != "xlsx":
        return "text/csv; charset=utf-8", _csv_bytes(), f"资产导出_{date}.csv"

    # Excel：优先 openpyxl（带样式 + 统计页），缺失则降级 CSV
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception:
        return "text/csv; charset=utf-8", _csv_bytes(), f"资产导出_{date}.csv"

    header_fill = PatternFill("solid", fgColor="005FB8")
    header_font = Font(color="FFFFFF", bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    warn_fill = PatternFill("solid", fgColor="FFC7CE")
    warn_font = Font(color="9C0006")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "资产列表"
    for ci, k in enumerate(keys, 1):
        c = ws.cell(row=1, column=ci, value=k)
        c.fill = header_fill
        c.font = header_font
        c.alignment = header_align
        c.border = border

    health_idx = keys.index("硬盘健康") + 1 if "硬盘健康" in keys else None
    for ri, m in enumerate(machines, 2):
        for ci, k in enumerate(keys, 1):
            c = ws.cell(row=ri, column=ci, value=str(m.get(k, "")))
            c.alignment = Alignment(vertical="center")
            c.border = border
        if health_idx is not None:
            hv = str(m.get("硬盘健康", ""))
            if any(x in hv for x in ("异常", "警告", "失败", "损坏", "坏")):
                for ci in range(1, len(keys) + 1):
                    cc = ws.cell(row=ri, column=ci)
                    cc.fill = warn_fill
                    cc.font = warn_font

    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(keys))
    ws.auto_filter.ref = f"A1:{last_col}{max(1, len(machines) + 1)}"

    for ci, k in enumerate(keys, 1):
        ml = len(str(k))
        for ri in range(2, len(machines) + 2):
            v = ws.cell(row=ri, column=ci).value
            if v is not None:
                w = sum(2 if ord(ch) > 127 else 1 for ch in str(v))
                ml = max(ml, w)
        ws.column_dimensions[get_column_letter(ci)].width = max(8, min(46, ml + 2))

    # 统计页：按 品牌 / 设备类型 / 部门
    stats = wb.create_sheet("统计")

    def count_by(key):
        d = {}
        for m in machines:
            v = str(m.get(key, "未填写") or "未填写").strip() or "未填写"
            d[v] = d.get(v, 0) + 1
        return sorted(d.items(), key=lambda kv: (-kv[1], kv[0]))

    row = 1
    for title, key in (("按品牌统计", "品牌"),
                       ("按设备类型统计", "设备类型"),
                       ("按部门统计", "部门")):
        stats.cell(row=row, column=1, value=title).font = Font(bold=True, size=12, color="005FB8")
        row += 1
        for j, h in enumerate(("类别", "数量"), 1):
            c = stats.cell(row=row, column=j, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = header_align
            c.border = border
        row += 1
        for name, cnt in count_by(key):
            a = stats.cell(row=row, column=1, value=name)
            b = stats.cell(row=row, column=2, value=cnt)
            a.border = border
            b.border = border
            b.alignment = Alignment(horizontal="center")
            row += 1
        row += 1
    stats.column_dimensions["A"].width = 22
    stats.column_dimensions["B"].width = 10

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            out.getvalue(), f"资产导出_{date}.xlsx")


def render_dashboard(machines):
    cols = EXPORT_COLS
    rows = []
    for m in machines:
        tds = "".join(f"<td>{str(m.get(c, '')).replace('<', '&lt;').replace('>', '&gt;')}</td>" for _, c in cols)
        rows.append(f"<tr>{tds}</tr>")
    head = "".join(f"<th>{label}</th>" for label, _ in cols)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>资产采集仪表盘</title>
<style>
 body{{font-family:-apple-system,'Microsoft YaHei',sans-serif;margin:0;background:#f5f6f8;color:#222}}
 header{{background:#005fb8;color:#fff;padding:14px 24px;font-size:18px;font-weight:bold}}
 .meta{{padding:10px 24px;color:#666;font-size:13px;display:flex;align-items:center;gap:14px;flex-wrap:wrap}}
 .meta a.btn{{background:#005fb8;color:#fff;text-decoration:none;padding:6px 12px;border-radius:6px;font-size:13px}}
 .meta a.ghost{{background:#eef3f8;color:#005fb8;border:1px solid #cfe0f0}}
 table{{border-collapse:collapse;width:96%;margin:0 24px 30px;background:#fff;box-shadow:0 1px 3px rgba(0,0,0,.1)}}
 th,td{{border:1px solid #e3e6ea;padding:8px 10px;font-size:13px;text-align:left;max-width:260px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}}
 th{{background:#eef3f8;position:sticky;top:0}}
 tr:nth-child(even){{background:#fafbfc}}
 .ok{{color:#2e7d32;font-weight:bold}}
 .warn{{color:#e65100;font-weight:bold}}
 .bad{{color:#c62828;font-weight:bold}}
</style></head>
<body>
<header>🖥️ 计算机资产采集仪表盘</header>
<div class="meta">
  <span>共 {len(machines)} 台机器 · 刷新时间 {now}</span>
  <a class="btn" href="/export?fmt=xlsx">⬇ 导出 Excel</a>
  <a class="btn ghost" href="/export?fmt=csv">⬇ 导出 CSV</a>
  <a href="/">点此刷新</a>
</div>
<div style="overflow:auto">
<table>
<thead><tr>{head}</tr></thead>
<tbody>
{''.join(rows)}
</tbody>
</table>
</div>
</body></html>"""
    return html


class _Handler(BaseHTTPRequestHandler):
    def _send(self, code, body, content_type="application/json; charset=utf-8"):
        if isinstance(body, (dict, list)):
            body = json.dumps(body, ensure_ascii=False)
        data = body.encode("utf-8") if isinstance(body, str) else body
        self.send_response(code)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(data)

    def do_POST(self):
        if self.path.rstrip("/") == "/api/report":
            try:
                length = int(self.headers.get("Content-Length", 0))
                raw = self.rfile.read(length) if length else b""
                data = json.loads(raw.decode("utf-8"))
            except Exception as e:
                self._send(400, {"ok": False, "error": str(e)})
                return
            if not isinstance(data, dict):
                self._send(400, {"ok": False, "error": "body must be json object"})
                return
            path = save_report(data)
            self._send(200, {"ok": True, "saved": os.path.basename(path)})
        else:
            self._send(404, {"ok": False, "error": "not found"})

    def do_GET(self):
        p = urlparse(self.path)
        route = p.path.rstrip("/") or "/"
        if route in ("/api/machines", "/api/report"):
            self._send(200, list_reports())
        elif route == "/export":
            qs = parse_qs(p.query)
            fmt = (qs.get("fmt", ["xlsx"])[0] or "xlsx").lower()
            q = qs.get("q", [""])[0] if qs.get("q") else ""
            machines = apply_query(list_reports(), q)
            ctype, data, fname = build_export(machines, fmt)
            # HTTP 头只能 latin-1：filename 用 ASCII 兜底名，中文名放 filename*（RFC 5987）
            date = datetime.now().strftime("%Y%m%d")
            ext = "xlsx" if fmt == "xlsx" else "csv"
            ascii_name = f"asset_export_{date}.{ext}"
            disp = f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(fname)}"
            self.send_response(200)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Access-Control-Allow-Origin", "*")
            self.send_header("Content-Disposition", disp)
            self.end_headers()
            self.wfile.write(data)
        elif route in ("/", "/dashboard"):
            self._send(200, render_dashboard(list_reports()), "text/html; charset=utf-8")
        else:
            self._send(404, {"ok": False, "error": "not found"})

    def log_message(self, *args):
        pass  # 静默日志，避免刷屏


class BatchServer:
    """中心采集服务器（可被 GUI 启停，也可独立运行）"""

    def __init__(self, host="0.0.0.0", port=8000):
        self.host = host
        self.port = port
        self.httpd = None
        self._thread = None

    def start(self):
        if self.httpd:
            return False, "已在运行"
        try:
            self.httpd = ThreadingHTTPServer((self.host, self.port), _Handler)
        except OSError as e:
            return False, f"启动失败（端口可能被占用）：{e}"
        self._thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self._thread.start()
        return True, f"监听 http://{self.host}:{self.port}（本机访问 http://127.0.0.1:{self.port}）"

    def stop(self):
        if self.httpd:
            self.httpd.shutdown()
            self.httpd.server_close()
            self.httpd = None
            self._thread = None
            return True
        return False

    def is_running(self):
        return self.httpd is not None

    def url(self):
        return f"http://127.0.0.1:{self.port}"


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="计算机资产中心采集服务器")
    ap.add_argument("--host", default="0.0.0.0",
                    help="监听地址，0.0.0.0=允许局域网/公网访问（默认），127.0.0.1=仅本机")
    ap.add_argument("--port", type=int, default=8000, help="监听端口（默认 8000）")
    a = ap.parse_args()

    srv = BatchServer(host=a.host, port=a.port)
    ok, msg = srv.start()
    print(msg)
    if ok:
        # 局域网模式：提示本机内网地址，方便分发给其它电脑
        try:
            import socket
            s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            s.connect(("8.8.8.8", 80))
            lan = s.getsockname()[0]
            s.close()
            print(f"局域网其它电脑请访问：http://{lan}:{a.port}")
        except Exception:
            pass
        try:
            webbrowser.open(srv.url())
        except Exception:
            pass
        print("按 Ctrl+C 停止服务器...")
        try:
            while True:
                threading.Event().wait(1)
        except KeyboardInterrupt:
            srv.stop()
            print("服务器已停止")
